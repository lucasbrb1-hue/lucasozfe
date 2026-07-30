import os
import tempfile
import unittest

try:
    import docx  # noqa: F401

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

from spt_piles import depth_solver as ds
from spt_piles import pile_group as pg
from spt_piles.loads import FoundationLoad
from spt_piles.models import PileGeometry, SPTProfile
from spt_piles.reinforcement import design_reinforcement
from spt_piles.pile_group import compute_batch_reinforcement


def build_profile() -> SPTProfile:
    profile = SPTProfile()
    for depth in range(1, 26):
        profile.add(depth, 4 + depth, "areia")
    return profile


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@unittest.skipUnless(HAS_DOCX, "python-docx não instalado")
class TestBatchMemorial(unittest.TestCase):
    def test_build_batch_memorial_uniformized(self):
        from spt_piles.batch_memorial import build_batch_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        loads = [
            FoundationLoad("P1", 300.0, 1),
            FoundationLoad("B2", 900.0, 3),
            FoundationLoad("P3", 600.0, 1),
        ]
        designs = pg.compute_individual_designs(loads, profile, geometry, "pre_moldada", method=ds.METHOD_DQ)
        pg.apply_group_uniformization(designs, n_groups=2)
        reinforcement = design_reinforcement(geometry, axial_load_kn=600.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "batch.docx")
            build_batch_memorial(
                out_path,
                loads,
                designs,
                profile,
                geometry,
                "pre_moldada",
                ds.METHOD_DQ,
                safety_factor=2.0,
                reinforcement=reinforcement,
                uniformized=True,
                n_groups=2,
            )
            self.assertTrue(os.path.exists(out_path))
            doc = docx.Document(out_path)
            text = _all_text(doc)
            self.assertIn("Memorial de Cálculo em Lote", text)
            self.assertIn("P1", text)
            self.assertIn("B2", text)
            self.assertIn("Grupo 1", text)
            self.assertIn("Uniformização de profundidades", text)
            self.assertGreaterEqual(len(doc.tables), 4)

    def test_build_batch_memorial_not_uniformized_and_infeasible_pile(self):
        from spt_piles.batch_memorial import build_batch_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=20)
        loads = [
            FoundationLoad("P1", 300.0, 1),
            FoundationLoad("P2", 5_000_000.0, 1),
        ]
        designs = pg.compute_individual_designs(loads, profile, geometry, "escavada", method=ds.METHOD_AV)
        pg.apply_no_uniformization(designs)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "batch.docx")
            build_batch_memorial(
                out_path,
                loads,
                designs,
                profile,
                geometry,
                "escavada",
                ds.METHOD_AV,
                safety_factor=2.0,
                reinforcement=None,
                uniformized=False,
            )
            doc = docx.Document(out_path)
            text = _all_text(doc)
            self.assertIn("INVIÁVEL", text)
            self.assertIn("estaca(s) não atingem", text)

    def test_partial_armor_length_reflected_per_pile_in_table(self):
        from spt_piles.reinforcement import effective_armor_length_m
        from spt_piles.batch_memorial import build_batch_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        loads = [
            FoundationLoad("P1", 300.0, 1),
            FoundationLoad("P2", 700.0, 1),
        ]
        designs = pg.compute_individual_designs(loads, profile, geometry, "pre_moldada", method=ds.METHOD_DQ)
        pg.apply_no_uniformization(designs)
        for d in designs:
            self.assertIsNotNone(d.adopted_depth_m)
        reinforcement = design_reinforcement(geometry, axial_load_kn=700.0, armor_length_m=5.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "batch.docx")
            build_batch_memorial(
                out_path, loads, designs, profile, geometry, "pre_moldada", ds.METHOD_DQ,
                safety_factor=2.0, reinforcement=reinforcement, uniformized=False,
            )
            text = _all_text(docx.Document(out_path))
            for d in designs:
                expected = f"{effective_armor_length_m(reinforcement, d.adopted_depth_m):.2f}"
                self.assertIn(expected, text)
            self.assertIn("Compr. armadura (m)", text)

    def test_batch_memorial_reports_structural_design_when_moment_given(self):
        from spt_piles.batch_memorial import build_batch_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=50)
        loads = [
            FoundationLoad("P1", 300.0, moment_kn_m=10.0),
            FoundationLoad("P2", 700.0, moment_kn_m=150.0, shear_kn=60.0),
        ]
        designs = pg.compute_individual_designs(loads, profile, geometry, "pre_moldada", method=ds.METHOD_DQ)
        pg.apply_no_uniformization(designs)
        reinforcement = compute_batch_reinforcement(loads, geometry)
        self.assertIsNotNone(reinforcement.structural)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "batch.docx")
            build_batch_memorial(
                out_path, loads, designs, profile, geometry, "pre_moldada", ds.METHOD_DQ,
                safety_factor=2.0, reinforcement=reinforcement, uniformized=False,
            )
            text = _all_text(docx.Document(out_path))
            self.assertIn("Dimensionamento estrutural", text)
            self.assertIn("P2", text)


if __name__ == "__main__":
    unittest.main()
