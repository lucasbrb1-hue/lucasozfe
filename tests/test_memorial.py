import os
import tempfile
import unittest

try:
    import docx  # noqa: F401

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

from spt_piles import depth_solver as ds
from spt_piles.models import PileGeometry, SPTProfile
from spt_piles.reinforcement import design_reinforcement


def build_profile() -> SPTProfile:
    profile = SPTProfile()
    for depth in range(1, 16):
        profile.add(depth, 5 + depth, "areia")
    return profile


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@unittest.skipUnless(HAS_DOCX, "python-docx não instalado")
class TestMemorial(unittest.TestCase):
    def test_build_memorial_creates_docx_with_expected_sections(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        solver_result = ds.solve(profile, geometry, "pre_moldada", load_kn=300.0, method=ds.METHOD_BOTH)
        reinforcement = design_reinforcement(geometry, axial_load_kn=300.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(
                out_path,
                profile,
                geometry,
                "pre_moldada",
                solver_result,
                reinforcement,
                safety_factor=2.0,
                water_table_depth_m=3.5,
                water_table_found=True,
            )
            self.assertTrue(os.path.exists(out_path))

            doc = docx.Document(out_path)
            full_text = _all_text(doc)
            self.assertIn("Memorial de Cálculo", full_text)
            self.assertIn("Décourt-Quaresma", full_text)
            self.assertIn("Aoki-Velloso", full_text)
            self.assertIn("3.50 m", full_text)
            self.assertGreaterEqual(len(doc.tables), 4)

    def test_build_memorial_single_method_only(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        solver_result = ds.solve(profile, geometry, "escavada", load_kn=300.0, method=ds.METHOD_DQ)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(
                out_path, profile, geometry, "escavada", solver_result, None, safety_factor=2.0,
            )
            doc = docx.Document(out_path)
            full_text = _all_text(doc)
            self.assertIn("Décourt-Quaresma", full_text)
            self.assertNotIn("rp = (K · Np) / F1", full_text)

    def test_build_memorial_reports_full_length_armor_by_default(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        solver_result = ds.solve(profile, geometry, "pre_moldada", load_kn=300.0, method=ds.METHOD_DQ)
        reinforcement = design_reinforcement(geometry, axial_load_kn=300.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(out_path, profile, geometry, "pre_moldada", solver_result, reinforcement, safety_factor=2.0)
            full_text = _all_text(docx.Document(out_path))
            self.assertIn("toda a profundidade da estaca", full_text)

    def test_build_memorial_reports_partial_armor_length_and_effective_value(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        solver_result = ds.solve(profile, geometry, "pre_moldada", load_kn=300.0, method=ds.METHOD_DQ)
        self.assertIsNotNone(solver_result.required_depth_m)
        reinforcement = design_reinforcement(geometry, axial_load_kn=300.0, armor_length_m=6.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(out_path, profile, geometry, "pre_moldada", solver_result, reinforcement, safety_factor=2.0)
            full_text = _all_text(docx.Document(out_path))
            self.assertIn("limitada a 6.00 m", full_text)
            self.assertIn("comprimento efetivo de armadura", full_text)

    def test_build_memorial_reports_structural_design_when_moment_given(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=50)
        solver_result = ds.solve(profile, geometry, "pre_moldada", load_kn=600.0, method=ds.METHOD_DQ)
        reinforcement = design_reinforcement(geometry, axial_load_kn=600.0, moment_kn_m=80.0, shear_kn=60.0)
        self.assertIsNotNone(reinforcement.structural)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(out_path, profile, geometry, "pre_moldada", solver_result, reinforcement, safety_factor=2.0)
            full_text = _all_text(docx.Document(out_path))
            self.assertIn("Dimensionamento estrutural", full_text)
            self.assertIn("Força normal de cálculo (Nd)", full_text)
            self.assertIn("Força cortante de cálculo (Vd)", full_text)
            self.assertIn(f"{reinforcement.structural.flexo_check.utilization * 100:.0f} %", full_text)

    def test_build_memorial_without_moment_has_no_structural_section(self):
        from spt_piles.memorial import build_memorial

        profile = build_profile()
        geometry = PileGeometry(diameter_cm=40)
        solver_result = ds.solve(profile, geometry, "pre_moldada", load_kn=300.0, method=ds.METHOD_DQ)
        reinforcement = design_reinforcement(geometry, axial_load_kn=300.0)

        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "memorial.docx")
            build_memorial(out_path, profile, geometry, "pre_moldada", solver_result, reinforcement, safety_factor=2.0)
            full_text = _all_text(docx.Document(out_path))
            self.assertNotIn("Dimensionamento estrutural", full_text)


if __name__ == "__main__":
    unittest.main()
