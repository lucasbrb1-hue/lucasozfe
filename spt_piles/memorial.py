"""Geração do memorial de cálculo completo (.docx) do dimensionamento da estaca.

Reúne, em um único documento Word: dados de entrada (perfil de SPT e, se
usada, a proveniência da extração por IA), a metodologia com as fórmulas
empregadas, a memória de cálculo passo a passo (valores intermediários) para
cada método selecionado, o resultado adotado e a armação sugerida.
"""

from __future__ import annotations

from datetime import datetime

from . import aoki_velloso, decourt_quaresma
from .ai_extraction import ExtractedSPTReport
from .depth_solver import METHOD_AV, METHOD_BOTH, METHOD_DQ, DepthSolverResult
from .models import PileGeometry, SPTProfile
from .pile_factors import PILE_TYPES
from .reinforcement import ReinforcementResult
from .soil_data import get_soil

DISCLAIMER = (
    "Este memorial é gerado por uma ferramenta de apoio ao pré-dimensionamento "
    "geotécnico/estrutural, com base em métodos semi-empíricos (Décourt-Quaresma "
    "e/ou Aoki-Velloso) e coeficientes de referência da literatura técnica. Os "
    "resultados NÃO substituem a investigação geotécnica completa, a análise "
    "crítica e a Anotação/Registro de Responsabilidade Técnica (ART/RRT) de um "
    "engenheiro habilitado. Validar sempre com as normas vigentes (NBR 6118, "
    "NBR 6122 e demais aplicáveis) antes de qualquer uso em projeto executivo "
    "ou execução de obra."
)


def _add_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        pass
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def build_memorial(
    output_path: str,
    profile: SPTProfile,
    geometry: PileGeometry,
    pile_type: str,
    solver_result: DepthSolverResult,
    reinforcement: ReinforcementResult | None,
    safety_factor: float,
    water_table_depth_m: float | None = None,
    water_table_found: bool = False,
    ai_extraction: ExtractedSPTReport | None = None,
) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    doc.add_heading("Memorial de Cálculo - Estaca dimensionada via SPT", level=0)
    doc.add_paragraph(f"Data de emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    doc.add_heading("1. Dados gerais", level=1)
    _add_table(
        doc,
        ["Parâmetro", "Valor"],
        [
            ["Tipo de estaca", PILE_TYPES.get(pile_type, pile_type)],
            ["Diâmetro", f"{geometry.diameter_cm:.1f} cm"],
            ["Carga de projeto", f"{solver_result.load_kn:.1f} kN"],
            ["Fator de segurança global adotado", f"{safety_factor:.2f}"],
            ["Método(s) de cálculo", _method_label(solver_result.method)],
            [
                "Nível d'água (N.A.)",
                (f"{water_table_depth_m:.2f} m" if water_table_found and water_table_depth_m is not None
                 else ("identificado, profundidade não especificada" if water_table_found else "não identificado / seco")),
            ],
        ],
    )

    if ai_extraction is not None:
        doc.add_heading("2. Proveniência dos dados (extração por IA)", level=1)
        p = doc.add_paragraph()
        p.add_run(
            "Os dados de sondagem abaixo foram inicialmente extraídos automaticamente "
            "por um modelo de IA (Claude, Anthropic) a partir do PDF do laudo SPT, e "
            "revisados/confirmados manualmente pelo usuário antes de serem utilizados "
            "neste cálculo."
        ).italic = True
        doc.add_paragraph(f"Arquivo de origem: {ai_extraction.source_pdf_path}")
        doc.add_paragraph(f"Modelo de IA utilizado: {ai_extraction.model_used}")
        if ai_extraction.borehole_id:
            doc.add_paragraph(f"Identificação do furo: {ai_extraction.borehole_id}")
        if ai_extraction.notes:
            doc.add_paragraph(f"Observações do laudo: {ai_extraction.notes}")

    section_n = 3 if ai_extraction is not None else 2
    doc.add_heading(f"{section_n}. Perfil de sondagem SPT (dados de entrada)", level=1)
    _add_table(
        doc,
        ["Profundidade (m)", "N-SPT", "Tipo de solo"],
        [[f"{p.depth_m:.2f}", p.n_spt, get_soil(p.soil_key).label] for p in profile.points],
    )
    section_n += 1

    doc.add_heading(f"{section_n}. Metodologia", level=1)
    _add_methodology(doc, solver_result.method)
    section_n += 1

    doc.add_heading(f"{section_n}. Memória de cálculo - capacidade de carga por profundidade", level=1)
    if solver_result.method in (METHOD_DQ, METHOD_BOTH):
        doc.add_heading("Décourt-Quaresma", level=2)
        _add_dq_table(doc, profile, geometry, pile_type, safety_factor)
    if solver_result.method in (METHOD_AV, METHOD_BOTH):
        doc.add_heading("Aoki-Velloso", level=2)
        _add_av_table(doc, profile, geometry, pile_type, safety_factor)
    section_n += 1

    doc.add_heading(f"{section_n}. Resultado adotado", level=1)
    if solver_result.required_depth_m is not None:
        doc.add_paragraph(
            f"Profundidade mínima necessária: {solver_result.required_depth_m:.2f} m, "
            f"para atender Qadm ≥ {solver_result.load_kn:.1f} kN "
            f"({_method_label(solver_result.method)})."
        )
    else:
        doc.add_paragraph(
            "Nenhuma profundidade dentro do perfil de SPT informado atinge a carga de "
            "projeto. É necessário aumentar o comprimento da sondagem, revisar o "
            "diâmetro da estaca ou a carga de projeto."
        )
    section_n += 1

    if reinforcement is not None:
        doc.add_heading(f"{section_n}. Armação sugerida", level=1)
        _add_reinforcement_section(doc, reinforcement, pile_depth_m=solver_result.required_depth_m)
        section_n += 1

    doc.add_heading(f"{section_n}. Aviso", level=1)
    warn = doc.add_paragraph(DISCLAIMER)
    for run in warn.runs:
        run.italic = True
        run.font.size = Pt(9)

    doc.save(output_path)


def _method_label(method: str) -> str:
    return {
        METHOD_DQ: "Décourt-Quaresma",
        METHOD_AV: "Aoki-Velloso",
        METHOD_BOTH: "Décourt-Quaresma e Aoki-Velloso (governa o mais conservador)",
    }.get(method, method)


def _add_methodology(doc, method: str) -> None:
    if method in (METHOD_DQ, METHOD_BOTH):
        doc.add_paragraph("Método de Décourt-Quaresma (1978, fatores de Décourt 1996):", style=None)
        doc.add_paragraph("rp = α · K · Np      (resistência de ponta, kPa)")
        doc.add_paragraph("rl = β · 10 · (Nl/3 + 1)      (resistência lateral, kPa)")
        doc.add_paragraph("Qp = rp · Ap ;  Ql = rl · U · L ;  Qu = Qp + Ql ;  Qadm = Qu / FS")
        doc.add_paragraph(
            "Np = média do N-SPT na região da ponta (leitura na cota, uma acima e uma "
            "abaixo). Nl = média do N-SPT ao longo do fuste (cada leitura limitada ao "
            "intervalo [3, 50]). K = coeficiente de solo (kPa). α, β = fatores empíricos "
            "por tipo de estaca e tipo de solo."
        )
    if method in (METHOD_AV, METHOD_BOTH):
        doc.add_paragraph("Método de Aoki-Velloso (1975, com fatores F1/F2 revisados):")
        doc.add_paragraph("rp = (K · Np) / F1      (resistência de ponta, kPa)")
        doc.add_paragraph("rl,i = (α/100 · K · Ni) / F2      (resistência lateral da camada i, kPa)")
        doc.add_paragraph("Qp = rp · Ap ;  Ql = U · Σ(rl,i · Δli) ;  Qu = Qp + Ql ;  Qadm = Qu / FS")
        doc.add_paragraph(
            "K e α são coeficientes tabelados por tipo de solo. F1 e F2 são fatores "
            "empíricos por tipo de estaca. Cada leitura de SPT é tratada como "
            "representativa de uma camada de espessura igual à média das distâncias "
            "às leituras vizinhas."
        )
    doc.add_paragraph(
        "Em ambos os métodos: Ap = área da ponta da estaca; U = perímetro da seção; "
        "L = profundidade/comprimento embutido; FS = fator de segurança global."
    )


def _add_dq_table(doc, profile: SPTProfile, geometry: PileGeometry, pile_type: str, fs: float) -> None:
    results = decourt_quaresma.calc_profile(profile, geometry, pile_type, safety_factor=fs, min_depth_m=1.0)
    headers = ["L (m)", "Np", "Nl", "K (kPa)", "α", "β", "Qp (kN)", "Ql (kN)", "Qu (kN)", "Qadm (kN)"]
    rows = [
        [
            f"{r.depth_m:.2f}",
            f"{r.np_tip:.1f}",
            f"{r.nl_shaft:.1f}",
            f"{r.k_kpa:.0f}",
            f"{r.alpha:.2f}",
            f"{r.beta:.2f}",
            f"{r.qp_kn:.1f}",
            f"{r.ql_kn:.1f}",
            f"{r.qu_kn:.1f}",
            f"{r.qadm_kn:.1f}",
        ]
        for r in results
    ]
    _add_table(doc, headers, rows)


def _add_av_table(doc, profile: SPTProfile, geometry: PileGeometry, pile_type: str, fs: float) -> None:
    results = aoki_velloso.calc_profile(profile, geometry, pile_type, safety_factor=fs, min_depth_m=1.0)
    headers = ["L (m)", "N ponta", "K ponta (MPa)", "F1", "F2", "Qp (kN)", "Ql (kN)", "Qu (kN)", "Qadm (kN)"]
    rows = [
        [
            f"{r.depth_m:.2f}",
            f"{r.n_tip:.0f}",
            f"{r.k_tip_mpa:.2f}",
            f"{r.f1:.2f}",
            f"{r.f2:.2f}",
            f"{r.qp_kn:.1f}",
            f"{r.ql_kn:.1f}",
            f"{r.qu_kn:.1f}",
            f"{r.qadm_kn:.1f}",
        ]
        for r in results
    ]
    _add_table(doc, headers, rows)


def _add_reinforcement_section(doc, reinforcement: ReinforcementResult, pile_depth_m: float | None = None) -> None:
    from .reinforcement import effective_armor_length_m

    doc.add_paragraph("As,min = ρmin · Ag")
    _add_table(
        doc,
        ["Parâmetro", "Valor"],
        [
            ["Área bruta da seção (Ag)", f"{reinforcement.gross_area_cm2:.1f} cm²"],
            ["Taxa mínima de armadura (ρmin)", f"{reinforcement.rho_min_pct:.2f} %"],
            ["Área de aço mínima (As,min)", f"{reinforcement.as_min_cm2:.2f} cm²"],
        ],
    )
    if reinforcement.longitudinal:
        lg = reinforcement.longitudinal
        doc.add_paragraph(
            f"Armadura longitudinal sugerida: {lg.n_bars} barras de φ{lg.bar_diameter_mm:.1f} mm "
            f"(As fornecido = {lg.as_provided_cm2:.2f} cm² ≥ As,min; espaçamento livre estimado "
            f"entre barras ≈ {lg.clear_spacing_cm:.1f} cm)."
        )
    else:
        doc.add_paragraph("Não foi encontrada combinação padrão viável de barras dentro dos limites adotados.")
    doc.add_paragraph(
        f"Estribos: φ{reinforcement.stirrup_diameter_mm:.1f} mm a cada "
        f"{reinforcement.stirrup_spacing_body_cm:.0f} cm no corpo da estaca, e a cada "
        f"{reinforcement.stirrup_spacing_top_cm:.0f} cm na zona de confinamento (primeiros "
        f"{reinforcement.confinement_length_m:.2f} m a partir do topo)."
    )
    if reinforcement.requested_armor_length_m is None:
        doc.add_paragraph(
            "Profundidade de armação: armadura longitudinal estendida por toda a profundidade da estaca."
        )
    else:
        armor_line = (
            f"Profundidade de armação: armadura longitudinal limitada a "
            f"{reinforcement.requested_armor_length_m:.2f} m a partir do topo da estaca"
        )
        if pile_depth_m is not None:
            eff = effective_armor_length_m(reinforcement, pile_depth_m)
            armor_line += (
                f" (estaca com {pile_depth_m:.2f} m de profundidade total; comprimento efetivo de "
                f"armadura = {eff:.2f} m)."
            )
        else:
            armor_line += " (limitada à profundidade real de cada estaca, se ela for menor que esse valor)."
        doc.add_paragraph(armor_line)
    for w in reinforcement.warnings:
        doc.add_paragraph(f"Aviso: {w}")
