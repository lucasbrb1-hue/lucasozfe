"""Geração do memorial de cálculo em lote (.docx): esforços importados,
profundidade individual/adotada de cada estaca (com ou sem uniformização) e
a armação comum adotada para o conjunto.
"""

from __future__ import annotations

from datetime import datetime

from .ai_extraction import ExtractedLoadsReport
from .loads import FoundationLoad
from .memorial import DISCLAIMER, _add_methodology, _add_reinforcement_section, _add_table, _method_label
from .models import PileGeometry, SPTProfile
from .pile_factors import PILE_TYPES
from .pile_group import PileDesign
from .reinforcement import ReinforcementResult, effective_armor_length_m
from .soil_data import get_soil

BATCH_DISCLAIMER_EXTRA = (
    "Quando um bloco possui mais de uma estaca, a carga por estaca foi obtida "
    "dividindo a carga característica do bloco igualmente pelo número de "
    "estacas informado - essa premissa não considera excentricidade de carga "
    "nem momentos que gerem distribuição desigual entre as estacas de um "
    "mesmo bloco, o que deve ser verificado pelo engenheiro responsável em "
    "blocos com geometria assimétrica ou cargas excêntricas relevantes. As "
    "cargas usadas são as CARACTERÍSTICAS (de serviço), não as majoradas "
    "(ELU/Nd) - a comparação com a capacidade admissível do solo (Qadm) já "
    "incorpora o fator de segurança geotécnico."
)


def build_batch_memorial(
    output_path: str,
    loads: list[FoundationLoad],
    designs: list[PileDesign],
    profile: SPTProfile,
    geometry: PileGeometry,
    pile_type: str,
    method: str,
    safety_factor: float,
    reinforcement: ReinforcementResult | None,
    uniformized: bool,
    n_groups: int | None = None,
    water_table_depth_m: float | None = None,
    water_table_found: bool = False,
    ai_loads_extraction: ExtractedLoadsReport | None = None,
) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    doc.add_heading("Memorial de Cálculo em Lote - Estacas dimensionadas via SPT", level=0)
    doc.add_paragraph(f"Data de emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    n_infeasible = sum(1 for d in designs if d.individual_required_depth_m is None)

    doc.add_heading("1. Dados gerais", level=1)
    _add_table(
        doc,
        ["Parâmetro", "Valor"],
        [
            ["Tipo de estaca", PILE_TYPES.get(pile_type, pile_type)],
            ["Diâmetro", f"{geometry.diameter_cm:.1f} cm"],
            ["Fator de segurança global adotado", f"{safety_factor:.2f}"],
            ["Método(s) de cálculo", _method_label(method)],
            ["Número de elementos importados", str(len(loads))],
            ["Número de estacas inviáveis (carga não atingida)", str(n_infeasible)],
            ["Uniformização de profundidades", ("Sim" if uniformized else "Não")],
            *([["Número de grupos/profundidades padrão", str(n_groups)]] if uniformized and n_groups else []),
            [
                "Nível d'água (N.A.)",
                (f"{water_table_depth_m:.2f} m" if water_table_found and water_table_depth_m is not None
                 else ("identificado, profundidade não especificada" if water_table_found else "não identificado / seco")),
            ],
        ],
    )

    section_n = 2
    if ai_loads_extraction is not None:
        doc.add_heading(f"{section_n}. Proveniência dos esforços (extração por IA)", level=1)
        p = doc.add_paragraph()
        p.add_run(
            "Os esforços abaixo foram inicialmente extraídos automaticamente por um "
            "modelo de IA (Claude, Anthropic) a partir do PDF do relatório de cargas, "
            "e revisados/confirmados manualmente pelo usuário antes deste cálculo."
        ).italic = True
        doc.add_paragraph(f"Arquivo de origem: {ai_loads_extraction.source_pdf_path}")
        doc.add_paragraph(f"Modelo de IA utilizado: {ai_loads_extraction.model_used}")
        if ai_loads_extraction.notes:
            doc.add_paragraph(f"Observações do relatório: {ai_loads_extraction.notes}")
        section_n += 1

    doc.add_heading(f"{section_n}. Perfil de sondagem SPT (dados de entrada)", level=1)
    _add_table(
        doc,
        ["Profundidade (m)", "N-SPT", "Tipo de solo"],
        [[f"{p.depth_m:.2f}", p.n_spt, get_soil(p.soil_key).label] for p in profile.points],
    )
    section_n += 1

    doc.add_heading(f"{section_n}. Metodologia", level=1)
    _add_methodology(doc, method)
    section_n += 1

    doc.add_heading(f"{section_n}. Esforços de fundação importados", level=1)
    if any(ld.combinations for ld in loads):
        doc.add_paragraph(
            "Elementos marcados com \"envoltória\" abaixo tiveram uma tabela completa de "
            "combinações de carregamento importada (N, Mx, My, Vx, Vy concomitantes por "
            "combinação) - a coluna Mk/Hk não mostra um único valor porque a armadura foi "
            "verificada contra TODAS as combinações da envoltória de cada elemento; a "
            "combinação mais exigente (governante) de cada estaca está indicada na seção de "
            "armação adiante."
        )
    _add_table(
        doc,
        [
            "Elemento", "Carga característica (kN)", "Nº de estacas no bloco",
            "Carga por estaca (kN)", "Mk (kN·m)", "Hk (kN)",
        ],
        [
            [
                ld.element_id, f"{ld.characteristic_load_kn:.1f}", ld.n_piles, f"{ld.load_per_pile_kn:.1f}",
                (f"envoltória ({len(ld.combinations)} combinações)" if ld.combinations
                 else (f"{ld.moment_kn_m:.1f}" if ld.moment_kn_m is not None else "-")),
                (f"envoltória ({len(ld.combinations)} combinações)" if ld.combinations
                 else (f"{ld.shear_kn:.1f}" if ld.shear_kn is not None else "-")),
            ]
            for ld in loads
        ],
    )
    section_n += 1

    doc.add_heading(f"{section_n}. Resultado por estaca", level=1)
    if uniformized:
        doc.add_paragraph(
            "Profundidades uniformizadas: as estacas foram agrupadas por profundidade "
            "individual necessária, adotando-se a maior profundidade de cada grupo "
            "para todas as estacas daquele grupo (nunca inferior à necessidade "
            "individual de nenhuma estaca do grupo)."
        )
    def _armor_txt(d: PileDesign) -> str:
        if d.adopted_depth_m is None or reinforcement is None:
            return "-"
        return f"{effective_armor_length_m(reinforcement, d.adopted_depth_m):.2f}"

    _add_table(
        doc,
        [
            "Elemento", "Carga/estaca (kN)", "Prof. individual necessária (m)", "Grupo",
            "Prof. adotada (m)", "Compr. armadura (m)",
        ],
        [
            [
                d.element_id,
                f"{d.load_per_pile_kn:.1f}",
                (f"{d.individual_required_depth_m:.2f}" if d.individual_required_depth_m is not None else "INVIÁVEL"),
                (d.group_label or "-"),
                (f"{d.adopted_depth_m:.2f}" if d.adopted_depth_m is not None else "-"),
                _armor_txt(d),
            ]
            for d in designs
        ],
    )
    if n_infeasible:
        doc.add_paragraph(
            f"Atenção: {n_infeasible} estaca(s) não atingem a carga necessária dentro da "
            "profundidade sondada. É necessário aumentar o comprimento da sondagem, "
            "revisar o diâmetro/tipo de estaca ou a carga desses elementos."
        )
    section_n += 1

    if reinforcement is not None:
        doc.add_heading(f"{section_n}. Armação adotada (comum a todas as estacas)", level=1)
        if reinforcement.requested_armor_length_m is None:
            doc.add_paragraph(
                "A armação abaixo (bitola, quantidade de barras, estribos) é comum a todas as "
                "estacas (mesmo diâmetro em todo o conjunto); a armadura longitudinal corre por "
                "toda a profundidade adotada de cada estaca, já indicada na coluna \"Compr. "
                "armadura (m)\" da tabela da seção anterior."
            )
        else:
            doc.add_paragraph(
                "A armação abaixo (bitola, quantidade de barras, estribos) é comum a todas as "
                "estacas (mesmo diâmetro em todo o conjunto); o comprimento efetivo da armadura "
                "longitudinal de cada estaca está na coluna \"Compr. armadura (m)\" da tabela da "
                "seção anterior (armadura parcial, limitada a "
                f"{reinforcement.requested_armor_length_m:.2f} m a partir do topo, ou à "
                "profundidade da estaca, o que for menor)."
            )
        _add_reinforcement_section(doc, reinforcement)
        section_n += 1

    doc.add_heading(f"{section_n}. Aviso", level=1)
    warn = doc.add_paragraph(DISCLAIMER + " " + BATCH_DISCLAIMER_EXTRA)
    for run in warn.runs:
        run.italic = True
        run.font.size = Pt(9)

    doc.save(output_path)
